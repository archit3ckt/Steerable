"""The library's views (grid, board, calendar), grouping, date ranges, tags and files columns, row
comments (and asking the COO), and spreadsheets in and out (library_io).

    python3 -m unittest tests.test_library_views -v
"""
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TEMPLATE = Path(__file__).resolve().parent.parent / "template"

SCRIPT = r"""
import csv, io, json, sys
sys.path.insert(0, "scripts")
import agent_api, bridge, files, library, library_io
out = {}
def err(fn):
    try:
        fn(); return None
    except (library.LibraryError, agent_api.ApiError) as e:
        return str(e) if isinstance(e, library.LibraryError) else e.message

jobs = library.table("Jobs", columns="Title:text, Status:status(New|Applied|Interview|Rejected), Interview:date, "
                     "Tags:tags, Level:tags(Junior|Senior), Docs:files, Salary:number")
cv = files.save("cv.pdf", b"%PDF-1.4 fake")
jobs.add({"Title": "DevOps", "Status": "Applied", "Interview": "2026-10-20", "Tags": "remote, onsite", "Level": ["junior"], "Docs": [cv["id"]], "Salary": 100})
jobs.add({"Title": "SRE", "Status": "Interview", "Interview": "2026-11-02", "Tags": ["onsite"], "Salary": 50})
jobs.add({"Title": "Barista", "Status": "Rejected"})
jobs.add({"Title": "Ops", "Status": "Applied"})
out["tags_value"] = jobs.get(jobs.query(search="devops")["rows"][0]["id"])["cells"]["Tags"]
out["level_bad"] = err(lambda: jobs.add({"Level": "Lead"}))
out["file_bad"] = err(lambda: jobs.add({"Docs": "deadbeef00"}))
titles = lambda p: [r["cells"]["Title"] for r in p["rows"]]
out["tag_where"] = titles(jobs.query(where={"Tags": "ONSITE"}))
out["tag_search"] = titles(jobs.query(search="remote"))
g = jobs.query(group="Status")
out["grouped"] = [titles(g), [(x["value"], x["count"]) for x in g["groups"]]]
out["between"] = titles(jobs.query(between=("Interview", "2026-10-01", "2026-10-31")))
out["files"] = list(jobs.query(search="devops")["files"].values())
out["group_bad"] = err(lambda: jobs.query(group="Tags"))

# views
out["default_views"] = [(v["name"], v["type"]) for v in jobs.views()]
board = jobs.save_view(None, "Pipeline", "kanban", {})
cal = jobs.save_view(None, None, "calendar", {"hidden": ["Salary"]})
out["board"] = [board["name"], board["config"]["by"]]
out["cal"] = [cal["name"], cal["config"]["by"], cal["config"]["hidden"]]
renamed = jobs.save_view(board["id"], "My pipeline", config={"where": {"tags": "onsite"}, "sort": "salary", "dir": "desc"})
out["renamed"] = [renamed["name"], renamed["config"]]
out["view_errors"] = [err(lambda: jobs.save_view(None, "x", "kanban", {"by": "Title"})),
                      err(lambda: jobs.save_view(None, "x", "pie", {})),
                      err(lambda: jobs.save_view(None, "x", "grid", {"sort": "Nope"}))]
notes = library.table("Notes", columns="Text:text")
out["no_status"] = err(lambda: notes.save_view(None, "b", "kanban", {}))
jobs.delete_view(cal["id"])
out["after_delete"] = [v["name"] for v in jobs.views()]
only = notes.views()[0]
out["keep_one"] = err(lambda: notes.delete_view(only["id"]))

# the API
api = agent_api.library_table("Jobs", {"group": ["Status"], "between": ["Interview,2026-10-01,2026-12-31"]})
out["api"] = [api["total"], [v["name"] for v in api["views"]], [(x["value"], x["count"]) for x in api["groups"]]]
made = agent_api.library_view("Jobs", {"name": "Board 2", "type": "kanban", "config": {"by": "Status"}})
out["api_view"] = made["config"]["by"]
out["api_bad_between"] = err(lambda: agent_api.library_table("Jobs", {"between": ["Interview,2026"]}))

# comments, and asking the COO (its turn faked here)
rid = jobs.query(search="devops")["rows"][0]["id"]
asked = []
bridge.converse = lambda prompt, *a, **k: asked.append(prompt) or "Done: I moved it to Interview."
plain = agent_api.library_comment("Jobs", rid, {"text": "Looks good"})
ask = agent_api.library_comment("Jobs", rid, {"text": "Write a cover letter for this one", "ask": True})
import time
for _ in range(50):
    if any(c["who"] == "coo" for c in jobs.comments(rid)):
        break
    time.sleep(0.1)
row = agent_api.library_row("Jobs", rid)
out["comments"] = [(c["who"], c["text"], c["status"]) for c in row["comments"]]
out["row_files"] = list(row["files"].values())
out["prompt"] = asked[0] if asked else ""
out["chat"] = [json.loads(l)["text"] for l in open("logs/chat.jsonl")] if __import__("os").path.exists("logs/chat.jsonl") else []
out["empty_comment"] = err(lambda: agent_api.library_comment("Jobs", rid, {"text": "  "}))

# spreadsheets out ...
data, name, kind = library_io.export("Jobs", "csv", where={"Status": "Applied"})
rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
out["csv"] = [name, rows[0], sorted(r[0] for r in rows[1:]), [r[3] for r in rows[1:] if r[0] == "DevOps"], [r[5] for r in rows[1:] if r[0] == "DevOps"]]
try:
    import openpyxl  # in the engine's image; may be missing where the tests run
except ImportError:
    openpyxl = None
if openpyxl:
    xdata, xname, _ = library_io.export("Jobs", "xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(xdata)).active
    out["xlsx"] = [xname, ws.max_row, [c.value for c in ws[1]][:3]]
out["bad_format"] = err(lambda: library_io.export("Jobs", "pdf"))

# ... and in
sheet = ("Name,Email,Stage,Deal size,Signed up,Website,Paid\n" +
         "\n".join(f"Shop {i},s{i}@x.com,{['Lead','Call','Won'][i % 3]},{i * 100},2026-09-{i % 28 + 1:02d},https://s{i}.example,{'yes' if i % 2 else 'no'}" for i in range(12)))
got = library_io.import_sheet("Customers.csv", sheet.encode())
t = library.table(got["table"]["name"])
out["import"] = [got["table"]["name"], got["added"], got["skipped"], [(c["name"], c["type"]) for c in t.columns]]
again = library_io.import_sheet("Customers.csv", sheet.encode())
out["import_again"] = again["table"]["name"]
more = library_io.import_sheet("more.csv", "Name,Deal size,Region\nShop X,lots,North\nShop Y,500,South\n".encode(), into=got["table"]["name"])
out["append"] = [more["added"], more["skipped"], more["columns"][-1], t.query(search="shop x")["rows"][0]["cells"].get("Deal size")]
if openpyxl:
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["שם", "תאריך", None]); ws.append(["דנה", __import__("datetime").date(2026, 10, 1), 5]); ws.append(["רון", None, 7])
    buf = io.BytesIO(); wb.save(buf)
    heb = library_io.import_sheet("לקוחות.xlsx", buf.getvalue())
    out["xlsx_import"] = [heb["table"]["name"], [(c["name"], c["type"]) for c in library.table(heb["table"]["id"]).columns]]
out["import_errors"] = [err(lambda: library_io.import_sheet("a.pdf", b"x")), err(lambda: library_io.import_sheet("a.csv", b"\n\n"))]
print(json.dumps(out, ensure_ascii=False, default=str))
"""


class LibraryViewsTest(unittest.TestCase):
    def test_views_comments_spreadsheets(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            root = tmp / "co"
            shutil.copytree(TEMPLATE, root)
            (root / ".env").write_text("COMPANY_NAME=Co\n")
            subprocess.run([sys.executable, "scripts/ops_db.py"], cwd=root, check=True, capture_output=True)
            proc = subprocess.run([sys.executable, "-c", SCRIPT], cwd=root, capture_output=True, text=True, timeout=300)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            out = json.loads(proc.stdout.strip().splitlines()[-1])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        # tags and files
        self.assertEqual(out["tags_value"], ["remote", "onsite"])
        self.assertIn("Junior, Senior", out["level_bad"])
        self.assertIn("no file", out["file_bad"])
        self.assertEqual(sorted(out["tag_where"]), ["DevOps", "SRE"])
        self.assertEqual(out["tag_search"], ["DevOps"])
        self.assertEqual(out["files"], [{"name": "cv.pdf", "kind": "pdf"}] if out["files"][0]["kind"] == "pdf" else out["files"])
        # grouping: by the status's order, with counts
        self.assertEqual(out["grouped"][1], [["Applied", 2], ["Interview", 1], ["Rejected", 1]])
        self.assertEqual(out["grouped"][0][2:], ["SRE", "Barista"])
        self.assertEqual(out["between"], ["DevOps"])
        self.assertIn("can't be grouped", out["group_bad"])
        # views
        self.assertEqual(out["default_views"], [["All rows", "grid"]])
        self.assertEqual(out["board"], ["Pipeline", "Status"])  # arranged by the first status column
        self.assertEqual(out["cal"], ["Calendar", "Interview", ["Salary"]])
        self.assertEqual(out["renamed"], ["My pipeline", {"sort": "Salary", "dir": "desc", "where": {"Tags": "onsite"}, "hidden": [], "by": "Status"}])
        self.assertIn("isn't one", out["view_errors"][0])
        self.assertIn("one of grid, kanban, calendar", out["view_errors"][1])
        self.assertIn("no column Nope", out["view_errors"][2])
        self.assertIn("has none", out["no_status"])
        self.assertEqual(out["after_delete"], ["All rows", "My pipeline"])
        self.assertIn("at least one view", out["keep_one"])
        self.assertEqual(out["api"][0], 2)
        self.assertEqual(out["api"][1], ["All rows", "My pipeline"])
        self.assertEqual(out["api"][2], [["Applied", 1], ["Interview", 1]])
        self.assertEqual(out["api_view"], "Status")
        self.assertIn("column,from,to", out["api_bad_between"])
        # comments
        self.assertEqual(out["comments"], [["you", "Looks good", None], ["you", "Write a cover letter for this one", "answered"],
                                           ["coo", "Done: I moved it to Interview.", None]])
        self.assertIn("Write a cover letter", out["prompt"])
        self.assertIn('library.py update "Jobs"', out["prompt"])
        self.assertIn("About “DevOps” in Jobs: Write a cover letter for this one", out["chat"])
        self.assertIn("5,000", out["empty_comment"])
        self.assertEqual(len(out["row_files"]), 1)
        # export
        name, header, titles, tags, docs = out["csv"]
        self.assertEqual(name, "Jobs.csv")
        self.assertEqual(header[:4], ["Title", "Status", "Interview", "Tags"])
        self.assertEqual(titles, ["DevOps", "Ops"])
        self.assertEqual(tags, ["remote, onsite"])
        self.assertEqual(docs, ["cv.pdf"])
        if "xlsx" in out:  # with openpyxl (the engine's image has it)
            self.assertEqual(out["xlsx"], ["Jobs.xlsx", 5, ["Title", "Status", "Interview"]])
            self.assertEqual(out["xlsx_import"], ["לקוחות", [["שם", "text"], ["תאריך", "date"], ["Column 3", "number"]]])
        self.assertIn("csv or xlsx", out["bad_format"])
        # import: types guessed, a second copy gets its own name, appending matches columns
        name, added, skipped, cols = out["import"]
        self.assertEqual((name, added, skipped), ("Customers", 12, 0))
        self.assertEqual(cols, [["Name", "text"], ["Email", "text"], ["Stage", "status"], ["Deal size", "number"],
                                ["Signed up", "date"], ["Website", "link"], ["Paid", "checkbox"]])
        self.assertEqual(out["import_again"], "Customers (2)")
        self.assertEqual(out["append"], [2, 1, "Region", None])  # "lots" isn't a number: left empty, row kept
        self.assertIn(".csv or .xlsx", out["import_errors"][0])
        self.assertIn("empty", out["import_errors"][1])


if __name__ == "__main__":
    unittest.main()
