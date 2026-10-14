"""Word, Excel and PowerPoint files attached in the chat get a text copy the COO can read
(scripts/files.py). Needs the engine's requirements (python-docx, openpyxl, python-pptx); skipped
without them.

    python3 -m unittest tests.test_files -v
"""
import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TEMPLATE = Path(__file__).resolve().parent.parent / "template"
HAVE_LIBS = all(importlib.util.find_spec(m) for m in ("docx", "openpyxl", "pptx"))

SCRIPT = """
import io, json, sys
sys.path.insert(0, "scripts")
import docx, openpyxl, pptx
import files

def saved(name, build):
    buf = io.BytesIO(); build(buf); return files.save(name, buf.getvalue())

def word(buf):
    d = docx.Document(); d.add_heading("Dana Levi", 0); d.add_paragraph("Senior bookkeeper, 8 years.")
    t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Skills"; t.rows[0].cells[1].text = "Xero, QuickBooks"
    d.save(buf)

def sheet(buf):
    b = openpyxl.Workbook(); s = b.active; s.title = "Prices"; s.append(["Item", "Price"]); s.append(["Pothos", 12.5]); b.save(buf)

def deck(buf):
    p = pptx.Presentation(); slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Q3 plan"; slide.notes_slide.notes_text_frame.text = "Mention the van"; p.save(buf)

out = {name: saved(name, fn) for name, fn in [("cv.docx", word), ("prices.xlsx", sheet), ("plan.pptx", deck)]}
out["broken"] = files.save("broken.docx", b"not really a word file")
out["texts"] = {k: open(v["text"]).read() for k, v in out.items() if v.get("text")}
out["prompt"] = files.for_prompt([out["cv.docx"], out["broken"]])
print(json.dumps(out))
"""


@unittest.skipUnless(HAVE_LIBS, "needs python-docx, openpyxl and python-pptx (template/requirements.txt)")
class OfficeTextTest(unittest.TestCase):
    def test_text_copies(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            root = tmp / "co"
            shutil.copytree(TEMPLATE, root)
            (root / ".env").write_text("COMPANY_NAME=Co\n")
            proc = subprocess.run([sys.executable, "-c", SCRIPT], cwd=root, capture_output=True, text=True, timeout=120)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            out = json.loads(proc.stdout)
            texts = out["texts"]
            self.assertIn("Dana Levi", texts["cv.docx"])
            self.assertIn("Xero, QuickBooks", texts["cv.docx"])  # tables too
            self.assertIn("## Sheet: Prices", texts["prices.xlsx"])
            self.assertIn("Pothos | 12.5", texts["prices.xlsx"])
            self.assertIn("Q3 plan", texts["plan.pptx"])
            self.assertIn("Notes: Mention the van", texts["plan.pptx"])
            # The COO is pointed to the text; a file that can't be read is named as such.
            self.assertIn(f"read its text in {out['cv.docx']['text']}", out["prompt"])
            self.assertNotIn("text", out["broken"])
            self.assertIn("a format you can't open directly", out["prompt"])
            # The original is still the file the owner downloads.
            self.assertTrue(out["cv.docx"]["path"].endswith("-cv.docx"))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
