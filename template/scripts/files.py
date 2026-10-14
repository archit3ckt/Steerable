"""Files the owner attaches in the chat: pictures, PDFs, documents.

Kept in the company folder under files/<year-month>/<id>-<name>, so they're exported and moved
with the company, and the COO opens them with Claude Code's Read tool, which reads pictures and
PDFs itself. Word, Excel and PowerPoint files it can't read, so a text copy is written next to
them (<id>.text.md) and the COO is pointed to that. A message names its files; the chat log
keeps their id, name, kind and size. Telegram attachments come here too (bridge.py, and the
cloud's shared bot through the agent API).

    python scripts/files.py list
"""
import json
import re
import secrets
import sys
from datetime import date
from pathlib import Path

from common import ROOT, log_run

DIR = ROOT / "files"
MAX_BYTES = 20 * 1024 * 1024
MAX_TOTAL = 1024 * 1024 * 1024  # all of a company's files together
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}
TYPES = {**IMAGES, ".pdf": "application/pdf", ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv",
         ".json": "application/json",
         ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
         ".epub": "application/epub+zip"}  # books the owner teaches the COO (books.py)
TEXT = {".txt", ".md", ".csv", ".json"}
ID = re.compile(r"[0-9a-f]{10}")
OFFICE = {".docx", ".xlsx", ".pptx"}
MAX_TEXT = 200_000  # characters of a document's text copy


class FileError(Exception):
    """Something to tell the owner, in plain words."""


def kind(name: str) -> str:
    ext = Path(name).suffix.lower()
    return "image" if ext in IMAGES else "pdf" if ext == ".pdf" else "text" if ext in TEXT else "other"


def _clean(name: str) -> str:
    name = Path(name.replace("\\", "/")).name.strip()
    stem, ext = Path(name).stem, Path(name).suffix.lower()
    stem = re.sub(r"[^\w\-. ]+", "", stem).strip(" .")[:80] or "file"
    return stem + re.sub(r"[^a-z0-9.]", "", ext)[:10]


def _used() -> int:
    return sum(p.stat().st_size for p in DIR.rglob("*") if p.is_file()) if DIR.exists() else 0


def describe(path: Path) -> dict:
    fid, _, name = path.name.partition("-")
    info = {"id": fid, "name": name, "kind": kind(name), "size": path.stat().st_size,
            "path": str(path.relative_to(ROOT))}
    if text_copy(path).exists():
        info["text"] = str(text_copy(path).relative_to(ROOT))
    return info


def save(name: str, data: bytes) -> dict:
    """Keep an attached file. Returns what the chat shows and the COO reads."""
    name = _clean(name)
    if not data:
        raise FileError("That file is empty.")
    if len(data) > MAX_BYTES:
        raise FileError("That file is too big. Attach files up to 20 MB.")
    if Path(name).suffix.lower() not in TYPES:
        raise FileError("I can read pictures (PNG, JPG, GIF, WebP), PDFs, text, CSV, Word, Excel or PowerPoint files and EPUB books.")
    if _used() + len(data) > MAX_TOTAL:
        raise FileError("Your files are full (1 GB). Remove some old ones first.")
    folder = DIR / date.today().strftime("%Y-%m")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{secrets.token_hex(5)}-{name}"
    path.write_bytes(data)
    if path.suffix.lower() in OFFICE:
        _write_text_copy(path)
    log_run("files", "saved", name=name, size=len(data))
    return describe(path)


def text_copy(path: Path) -> Path:
    return path.with_name(path.name.partition("-")[0] + ".text.md")


def _write_text_copy(path: Path) -> None:
    """The text of a Word, Excel or PowerPoint file, as Markdown, for the COO to read."""
    try:
        text = _office_text(path)
    except Exception as e:  # noqa: BLE001 - a damaged or odd file: the COO says it can't open it
        log_run("files", "no_text", name=path.name, error=repr(e)[:300])
        return
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT] + "\n\n(The rest is cut off: the document is very long.)"
    text_copy(path).write_text(f"# {path.name.partition('-')[2]}\n\n{text.strip()}\n")


def _office_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.append("\n".join(" | ".join(c.text.strip() for c in row.cells) for row in table.rows))
        return "\n\n".join(parts)
    if ext == ".xlsx":
        import openpyxl

        book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet in book.worksheets:
            rows = []
            for n, row in enumerate(sheet.iter_rows(values_only=True)):
                if n >= 5000:
                    rows.append("(more rows not shown)")
                    break
                if any(v is not None for v in row):
                    rows.append(" | ".join("" if v is None else str(v) for v in row))
            parts.append(f"## Sheet: {sheet.title}\n\n" + "\n".join(rows))
        return "\n\n".join(parts)
    import pptx

    deck = pptx.Presentation(str(path))
    parts = []
    for n, slide in enumerate(deck.slides, 1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        parts.append(f"## Slide {n}\n\n" + "\n".join(texts))
    return "\n\n".join(parts)


def find(fid: str) -> Path:
    if not ID.fullmatch(fid or ""):
        raise FileError("That file isn't here.")
    found = next(DIR.glob(f"*/{fid}-*"), None) if DIR.exists() else None
    if not found or not found.is_file():
        raise FileError("That file isn't here any more.")
    return found


def content_type(path: Path) -> str:
    return TYPES.get(path.suffix.lower(), "application/octet-stream")


def for_prompt(attached: list[dict]) -> str:
    """How a message's files are named to the COO, so it opens them."""
    if not attached:
        return ""
    lines = [f"- {a['path']} (“{a['name']}”, {a['kind']})"
             + (f": read its text in {a['text']}" if a.get("text") else
                ": a format you can't open directly" if Path(a["name"]).suffix.lower() in OFFICE else "")
             for a in attached]
    return ("\n\nThe owner attached these files (in the company folder). Open each with the Read tool "
            "before you answer:\n" + "\n".join(lines))


if __name__ == "__main__":
    if sys.argv[1:] == ["list"]:
        print(json.dumps([describe(p) for p in sorted(DIR.rglob("*")) if p.is_file() and "-" in p.name], indent=1))
    else:
        print(__doc__)
